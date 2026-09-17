"""Email one track's workbook as a ranked digest, with the .xlsx attached.

`autopilot.py run` stops before the mailer in manual_apply_only mode, so this is
the piece that actually delivers. Credentials come from the environment, never
from the repo:

    GMAIL_ADDRESS        sender + recipient default
    GMAIL_APP_PASSWORD   16-char Google app password (NOT the account password)

    python ci/send_digest.py --xlsx out/ai/jobs.xlsx --label "ML & AI" --top 25
"""
import argparse
import datetime
import html
import os
import pathlib
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))
import core  # noqa: E402


def load_jobs(xlsx: pathlib.Path) -> list[dict]:
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    # The per-date jobs sheet is the widest one; "Summary" is a small side sheet.
    ws = max(wb.worksheets, key=lambda w: (w.max_row or 0) * (w.max_column or 0))
    rows = ws.iter_rows(values_only=True)
    try:
        hdr = [str(h) if h is not None else "" for h in next(rows)]
    except StopIteration:
        return []
    return [dict(zip(hdr, r)) for r in rows]


def pick(d: dict, *needles: str):
    for k, v in d.items():
        low = k.lower()
        if all(n in low for n in needles):
            return v
    return None


def score_of(job: dict) -> int:
    v = pick(job, "score")
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return -1


def build_html(label: str, jobs: list[dict], shown: list[dict], stamp: str) -> str:
    rows = []
    for j in shown:
        title = html.escape(str(pick(j, "title") or "?"))
        company = html.escape(str(pick(j, "company") or "?"))
        loc = html.escape(str(pick(j, "location") or ""))
        url = str(pick(j, "apply", "url") or pick(j, "url") or "")
        fit = html.escape(str(pick(j, "fit") or "")[:160])
        sc = score_of(j)
        colour = "#16a34a" if sc >= 75 else "#ca8a04" if sc >= 55 else "#6b7280"
        link = (f'<a href="{html.escape(url, quote=True)}" '
                f'style="color:#2563eb;text-decoration:none;">{title}</a>') if url else title
        rows.append(
            f'<tr>'
            f'<td style="padding:6px 10px;font-weight:700;color:{colour};'
            f'border-bottom:1px solid #eee;">{sc if sc >= 0 else "-"}</td>'
            f'<td style="padding:6px 10px;border-bottom:1px solid #eee;">{link}'
            f'<div style="color:#6b7280;font-size:12px;">{company}'
            f'{" &middot; " + loc if loc else ""}</div>'
            f'{f"<div style=color:#9ca3af;font-size:11px;>{fit}</div>" if fit else ""}'
            f'</td></tr>'
        )
    scored = [j for j in jobs if score_of(j) >= 0]
    strong = len([j for j in scored if score_of(j) >= 75])
    review = len([j for j in scored if 55 <= score_of(j) < 75])
    return (
        f'<div style="font-family:-apple-system,Segoe UI,sans-serif;max-width:720px;">'
        f'<h2 style="margin:0 0 4px 0;">{html.escape(label)} &mdash; {len(jobs)} jobs</h2>'
        f'<div style="color:#6b7280;font-size:13px;margin-bottom:14px;">'
        f'{stamp} &middot; {strong} strong (75+) &middot; {review} worth a review (55-74) '
        f'&middot; full workbook attached</div>'
        f'<table style="border-collapse:collapse;width:100%;font-size:14px;">'
        f'<tr><th style="text-align:left;padding:6px 10px;border-bottom:2px solid #ddd;">Score</th>'
        f'<th style="text-align:left;padding:6px 10px;border-bottom:2px solid #ddd;">Role</th></tr>'
        f'{"".join(rows) or "<tr><td colspan=2 style=padding:10px;>No scored jobs this run.</td></tr>"}'
        f'</table></div>'
    )


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--label", required=True)
    ap.add_argument("--to", default="")
    ap.add_argument("--top", type=int, default=25)
    ap.add_argument("--min-score", type=int, default=0)
    args = ap.parse_args()

    address = os.environ.get("GMAIL_ADDRESS", "").strip()
    app_pw = os.environ.get("GMAIL_APP_PASSWORD", "").strip()
    to = args.to.strip() or address
    if not (address and app_pw and to):
        print("GMAIL_ADDRESS / GMAIL_APP_PASSWORD not set — skipping digest email",
              file=sys.stderr)
        return 0

    xlsx = pathlib.Path(args.xlsx)
    if not xlsx.exists():
        print(f"no workbook at {xlsx} — nothing to send", file=sys.stderr)
        return 0

    jobs = load_jobs(xlsx)
    ranked = sorted((j for j in jobs if score_of(j) >= args.min_score),
                    key=score_of, reverse=True)[:args.top]
    stamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M")

    # Attach under a dated name so saved mail doesn't collapse into one file.
    dated = xlsx.with_name(
        f"{args.label.replace(' & ', '_').replace(' ', '_')}"
        f"_{datetime.datetime.now().strftime('%Y-%m-%d_%Hh%M')}.xlsx"
    )
    dated.write_bytes(xlsx.read_bytes())

    mailer = core.Gmailer(address=address, app_password=app_pw, send=True)
    subject = f"{args.label}: {len(jobs)} jobs ({stamp})"
    mailer.send(to, subject, build_html(args.label, jobs, ranked, stamp), [dated])
    print(f"digest sent to {to}: {len(jobs)} jobs, {len(ranked)} listed, "
          f"attachment={dated.name}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
